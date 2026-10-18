"""FileOrbit's on-demand conversion process. SPDX-License-Identifier: AGPL-3.0-only."""
from __future__ import annotations

import configparser
import io
import json
import os
import shutil
import subprocess
import sys
import tarfile
import tempfile
import threading
import zipfile
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk


APP_DIR = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
SETTINGS_PATH = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "FileOrbit" / "settings.ini"
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
VIDEO_EXTS = {".mp4", ".mkv", ".mov", ".avi", ".webm"}
AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".aac"}
ARCHIVE_EXTS = {".zip", ".7z", ".tar", ".gz", ".tgz", ".xz"}


def lang(cfg: configparser.ConfigParser) -> str:
    return cfg.get("general", "language", fallback="ar")


def label(cfg: configparser.ConfigParser, ar: str, en: str) -> str:
    return en if lang(cfg) == "en" else ar


def load_settings() -> configparser.ConfigParser:
    cfg = configparser.ConfigParser()
    cfg["general"] = {
        "language": "ar", "right_hold": "1", "shift_drag": "1", "hold_ms": "500",
        "startup": "1", "review_text": "1", "review_images": "1",
        "sort": "name", "sort_desc": "0", "quality": "balanced",
    }
    if SETTINGS_PATH.exists():
        # The tray process uses a flat INI file; accept both that and sectioned INI.
        contents = SETTINGS_PATH.read_text(encoding="utf-8-sig")
        if "[" not in contents:
            contents = "[general]\n" + contents
        cfg.read_string(contents)
    return cfg


def save_settings(cfg: configparser.ConfigParser) -> None:
    SETTINGS_PATH.parent.mkdir(parents=True, exist_ok=True)
    with SETTINGS_PATH.open("w", encoding="utf-8") as out:
        for key, value in cfg["general"].items():
            out.write(f"{key}={value}\n")
    # The user explicitly controls this in Settings; no registry change happens elsewhere.
    try:
        import winreg
        key = winreg.CreateKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Windows\CurrentVersion\Run")
        if cfg.getboolean("general", "startup", fallback=True):
            tray_exe = APP_DIR / "FileOrbit.exe"
            if tray_exe.exists():
                winreg.SetValueEx(key, "FileOrbit", 0, winreg.REG_SZ, f'"{tray_exe}"')
        else:
            try:
                winreg.DeleteValue(key, "FileOrbit")
            except FileNotFoundError:
                pass
        winreg.CloseKey(key)
    except OSError:
        pass


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    for n in range(2, 10000):
        candidate = path.with_name(f"{path.stem} ({n}){path.suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError("Too many files with the same name")


def unique_dir(path: Path) -> Path:
    if not path.exists():
        return path
    for n in range(2, 10000):
        candidate = path.with_name(f"{path.name} ({n})")
        if not candidate.exists():
            return candidate
    raise RuntimeError("Too many folders with the same name")


def category(path: Path) -> str:
    if path.is_dir():
        return "folder"
    suffix = path.suffix.lower()
    if suffix in IMAGE_EXTS:
        return "image"
    if suffix == ".pdf":
        return "pdf"
    if suffix in AUDIO_EXTS:
        return "audio"
    if suffix in VIDEO_EXTS:
        return "video"
    if suffix in ARCHIVE_EXTS:
        return "archive"
    if suffix in {".txt", ".md", ".docx"}:
        return "document"
    return "other"


def capture_time(path: Path) -> float:
    try:
        from PIL import Image
        from datetime import datetime
        with Image.open(path) as image:
            exif = image.getexif()
            raw = exif.get(36867) or exif.get(306)
            if raw:
                return datetime.strptime(str(raw), "%Y:%m:%d %H:%M:%S").timestamp()
    except Exception:
        pass
    return path.stat().st_mtime


def sorted_images(files: list[Path], cfg: configparser.ConfigParser) -> list[Path]:
    mode = cfg.get("general", "sort", fallback="name")
    reverse = cfg.getboolean("general", "sort_desc", fallback=False)
    if mode == "time":
        return sorted(files, key=lambda p: (capture_time(p), p.name.casefold()), reverse=reverse)
    return sorted(files, key=lambda p: p.name.casefold(), reverse=reverse)


def pdf_from_images(files: list[Path], out: Path, cfg: configparser.ConfigParser) -> Path:
    from PIL import Image, ImageOps
    files = sorted_images(files, cfg)
    if cfg.getboolean("general", "review_images", fallback=True) and len(files) > 1:
        chosen = review_image_order(files, cfg)
        if chosen is None:
            raise InterruptedError("Cancelled")
        files = chosen
    pages = []
    try:
        for path in files:
            with Image.open(path) as original:
                image = ImageOps.exif_transpose(original).convert("RGB")
                pages.append(image.copy())
        out = unique_path(out)
        pages[0].save(out, "PDF", save_all=True, append_images=pages[1:], resolution=120.0)
        return out
    finally:
        for page in pages:
            page.close()


def pdf_to_images(path: Path, extension: str, dpi: int = 160) -> Path:
    import pymupdf as fitz
    folder = unique_dir(path.with_name(path.stem + "_pages"))
    folder.mkdir(parents=True)
    try:
        with fitz.open(path) as doc:
            for i, page in enumerate(doc, 1):
                pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
                pix.save(folder / f"{path.stem}_{i:03d}.{extension}")
        return folder
    except Exception:
        shutil.rmtree(folder, ignore_errors=True)
        raise


def tesseract_path() -> Path | None:
    local = APP_DIR / "assets" / "tesseract" / "tesseract.exe"
    if local.exists():
        return local
    system = shutil.which("tesseract")
    return Path(system) if system else None


def ocr_image_bytes(image: bytes, language: str = "ara+eng") -> str:
    binary = tesseract_path()
    if binary is None:
        raise RuntimeError("Offline OCR engine is missing. / محرك قراءة النص غير موجود")
    env = os.environ.copy()
    tessdata = APP_DIR / "assets" / "tesseract" / "tessdata"
    if tessdata.exists():
        env["TESSDATA_PREFIX"] = str(tessdata)
    proc = subprocess.run([str(binary), "stdin", "stdout", "-l", language, "--psm", "3"],
                          input=image, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env)
    if proc.returncode:
        raise RuntimeError(proc.stderr.decode("utf-8", errors="replace"))
    return proc.stdout.decode("utf-8", errors="replace")


def pdf_text(path: Path, force_ocr: bool = False) -> str:
    import pymupdf as fitz
    sections = []
    with fitz.open(path) as doc:
        for page in doc:
            text = page.get_text("text").strip()
            if force_ocr or len(text) < 15:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                text = ocr_image_bytes(pix.tobytes("png")).strip()
            sections.append(text)
    return "\n\n\f\n\n".join(sections)


def save_text_as(text: str, source: Path, target: str) -> Path:
    out = unique_path(source.with_suffix("." + target))
    if target == "txt":
        out.write_text(text, encoding="utf-8")
    elif target == "docx":
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        document = Document()
        for para in text.splitlines():
            p = document.add_paragraph(para)
            if any("\u0600" <= ch <= "\u06ff" for ch in para):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        document.save(out)
    else:
        raise ValueError(target)
    return out


def image_convert(path: Path, target: str, quality: int = 88) -> Path:
    from PIL import Image, ImageOps
    out = unique_path(path.with_suffix("." + target))
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source)
        if target in {"jpg", "jpeg", "bmp", "pdf"}:
            image = image.convert("RGB")
        image.save(out, quality=quality, optimize=target in {"jpg", "jpeg", "png"})
    return out


def pdf_compress(path: Path) -> Path:
    import pymupdf as fitz
    out = unique_path(path.with_name(path.stem + "_compressed.pdf"))
    with fitz.open(path) as doc:
        doc.save(out, garbage=4, deflate=True, clean=True)
    return out


def document_convert(path: Path, target: str) -> Path:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        text = "\n".join(p.text for p in Document(path).paragraphs)
    else:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
    if target in {"txt", "docx"}:
        return save_text_as(text, path, target)
    if target == "pdf":
        import html
        import pymupdf as fitz
        out = unique_path(path.with_suffix(".pdf"))
        lines = text.splitlines() or [""]
        paragraphs = []
        for line in lines:
            rtl = any("\u0590" <= ch <= "\u08ff" for ch in line)
            direction = "rtl" if rtl else "ltr"
            paragraphs.append(f'<p dir="{direction}">{html.escape(line) or "&nbsp;"}</p>')
        content = '<html><body>' + ''.join(paragraphs) + '</body></html>'
        css = 'body { font-family: sans-serif; font-size: 11pt; } p { margin: 0 0 8pt 0; }'
        story = fitz.Story(content, user_css=css)
        page_rect = fitz.paper_rect("a4")
        content_rect = page_rect + (48, 48, -48, -48)
        writer = fitz.DocumentWriter(str(out))
        try:
            more = 1
            while more:
                device = writer.begin_page(page_rect)
                more, _ = story.place(content_rect)
                story.draw(device)
                writer.end_page()
        finally:
            writer.close()
        return out
    raise ValueError(target)


def ffmpeg_exe() -> str:
    local = APP_DIR / "assets" / "ffmpeg" / "ffmpeg.exe"
    if local.exists():
        return str(local)
    system = shutil.which("ffmpeg")
    if system:
        return system
    raise RuntimeError("Offline FFmpeg engine is missing. / محرك الوسائط غير موجود")


def run_media(path: Path, target: str, mode: str, quality: str, root: tk.Tk) -> Path:
    out = unique_path(path.with_name(path.stem + ("_compressed" if mode == "compress" else "") + "." + target))
    cmd = [ffmpeg_exe(), "-hide_banner", "-loglevel", "error", "-y", "-i", str(path)]
    if mode == "extract_audio" or target in {"mp3", "wav", "flac", "m4a", "ogg", "aac"}:
        cmd += ["-vn"]
        if target == "mp3": cmd += ["-c:a", "libmp3lame", "-q:a", "3"]
        elif target == "m4a": cmd += ["-c:a", "aac", "-b:a", "192k"]
        elif target == "ogg": cmd += ["-c:a", "libvorbis", "-q:a", "5"]
    else:
        crf = {"high": "18", "balanced": "23", "small": "29"}.get(quality, "23")
        if target == "webm": cmd += ["-c:v", "libvpx-vp9", "-crf", crf, "-b:v", "0", "-c:a", "libopus"]
        else: cmd += ["-c:v", "libx264", "-preset", "fast", "-crf", crf, "-c:a", "aac", "-b:a", "160k"]
    cmd.append(str(out))
    progress = tk.Toplevel(root)
    progress.title("FileOrbit")
    progress.geometry("380x120")
    progress.resizable(False, False)
    ttk.Label(progress, text=path.name, wraplength=340).pack(pady=(15, 8))
    bar = ttk.Progressbar(progress, mode="indeterminate")
    bar.pack(fill="x", padx=20)
    bar.start(12)
    proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, creationflags=0x08000000)
    ttk.Button(progress, text="Cancel / إلغاء", command=proc.terminate).pack(pady=8)
    while proc.poll() is None:
        root.update()
        root.after(80)
    error = proc.stderr.read().decode("utf-8", errors="replace")
    progress.destroy()
    if proc.returncode:
        out.unlink(missing_ok=True)
        raise RuntimeError(error or "Media conversion was cancelled or failed")
    return out


def safe_member_path(folder: Path, name: str) -> Path:
    root = folder.resolve()
    candidate = (root / name).resolve()
    if candidate != root and root not in candidate.parents:
        raise ValueError("Archive contains a path outside the output folder")
    return candidate


def extract_archive(path: Path) -> Path:
    folder = unique_dir(path.with_name(path.name.split(".")[0]))
    folder.mkdir(parents=True)
    try:
        suffix = path.suffix.lower()
        if suffix == ".zip":
            with zipfile.ZipFile(path) as archive:
                for item in archive.infolist():
                    target = safe_member_path(folder, item.filename)
                    if item.is_dir(): target.mkdir(parents=True, exist_ok=True)
                    else:
                        target.parent.mkdir(parents=True, exist_ok=True)
                        with archive.open(item) as src, target.open("wb") as dst: shutil.copyfileobj(src, dst)
        elif suffix == ".7z":
            import py7zr
            with py7zr.SevenZipFile(path) as archive:
                for name in archive.getnames(): safe_member_path(folder, name)
                archive.extractall(folder)
        else:
            with tarfile.open(path, "r:*") as archive:
                for item in archive:
                    target = safe_member_path(folder, item.name)
                    if item.isdir(): target.mkdir(parents=True, exist_ok=True)
                    elif item.isfile():
                        target.parent.mkdir(parents=True, exist_ok=True)
                        src = archive.extractfile(item)
                        if src:
                            with src, target.open("wb") as dst: shutil.copyfileobj(src, dst)
                    # Links and device files are deliberately ignored.
        return folder
    except Exception:
        shutil.rmtree(folder, ignore_errors=True)
        raise


def create_archive(files: list[Path], target: str) -> Path:
    base = files[0].parent / (files[0].stem if len(files) == 1 else "FileOrbit Archive")
    extension = {"zip": ".zip", "7z": ".7z", "tar": ".tar", "tar.gz": ".tar.gz"}[target]
    out = unique_path(Path(str(base) + extension))
    try:
        if target == "zip":
            with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
                for path in files:
                    if path.is_dir():
                        for item in path.rglob("*"):
                            if item.is_file(): archive.write(item, arcname=str(Path(path.name) / item.relative_to(path)))
                    else: archive.write(path, arcname=path.name)
        elif target == "7z":
            import py7zr
            with py7zr.SevenZipFile(out, "w") as archive:
                for path in files:
                    if path.is_dir(): archive.writeall(path, arcname=path.name)
                    else: archive.write(path, arcname=path.name)
        else:
            with tarfile.open(out, "w:gz" if target == "tar.gz" else "w") as archive:
                for path in files: archive.add(path, arcname=path.name, recursive=True)
        return out
    except Exception:
        out.unlink(missing_ok=True)
        raise


ROOT: tk.Tk | None = None


def modal(title: str, width: int = 450, height: int = 320) -> tk.Toplevel:
    assert ROOT is not None
    window = tk.Toplevel(ROOT)
    window.title(title)
    window.geometry(f"{width}x{height}")
    window.minsize(width, height)
    window.transient(ROOT)
    window.grab_set()
    window.lift()
    return window


def choose_one(cfg: configparser.ConfigParser, title_ar: str, title_en: str,
               choices: list[tuple[str, str, str]], default: str | None = None) -> str | None:
    window = modal("FileOrbit", 420, 190)
    ttk.Label(window, text=label(cfg, title_ar, title_en), font=("Segoe UI", 12)).pack(pady=(22, 18))
    values = [label(cfg, ar, en) for _key, ar, en in choices]
    selected = tk.StringVar(value=values[next((i for i, item in enumerate(choices) if item[0] == default), 0)])
    combo = ttk.Combobox(window, values=values, textvariable=selected, state="readonly", width=32)
    combo.pack()
    result = [None]
    def finish():
        result[0] = choices[values.index(selected.get())][0]
        window.destroy()
    ttk.Button(window, text=label(cfg, "متابعة", "Continue"), command=finish).pack(pady=20)
    window.protocol("WM_DELETE_WINDOW", window.destroy)
    window.wait_window()
    return result[0]


def review_image_order(files: list[Path], cfg: configparser.ConfigParser) -> list[Path] | None:
    window = modal(label(cfg, "ترتيب الصور", "Image order"), 520, 480)
    ttk.Label(window, text=label(cfg, "رتّب الصور بسحبها أو استخدم الأزرار", "Arrange images before creating the PDF"),
              font=("Segoe UI", 11)).pack(pady=12)
    body = ttk.Frame(window)
    body.pack(fill="both", expand=True, padx=18)
    listing = tk.Listbox(body, font=("Segoe UI", 10), selectmode="browse")
    listing.pack(side="left", fill="both", expand=True)
    order = files[:]
    def refresh(selection: int = 0):
        listing.delete(0, "end")
        for i, path in enumerate(order, 1): listing.insert("end", f"{i:02d}  {path.name}")
        if order:
            selection = max(0, min(len(order) - 1, selection))
            listing.selection_set(selection)
            listing.see(selection)
    def move(delta: int):
        if not listing.curselection(): return
        index = listing.curselection()[0]
        other = max(0, min(len(order) - 1, index + delta))
        order[index], order[other] = order[other], order[index]
        refresh(other)
    buttons = ttk.Frame(body)
    buttons.pack(side="right", padx=10)
    ttk.Button(buttons, text="↑", width=5, command=lambda: move(-1)).pack(pady=5)
    ttk.Button(buttons, text="↓", width=5, command=lambda: move(1)).pack(pady=5)
    ttk.Button(buttons, text=label(cfg, "بالاسم", "By name"), command=lambda: resort("name")).pack(pady=5)
    ttk.Button(buttons, text=label(cfg, "بالوقت", "By time"), command=lambda: resort("time")).pack(pady=5)
    def resort(mode: str):
        order.sort(key=(lambda p: capture_time(p)) if mode == "time" else (lambda p: p.name.casefold()))
        refresh()
    # Dragging a list row reorders this conversion only; settings remain unchanged.
    drag_start = [None]
    def on_press(event): drag_start[0] = listing.nearest(event.y)
    def on_release(event):
        if drag_start[0] is None: return
        to = listing.nearest(event.y)
        if 0 <= drag_start[0] < len(order) and 0 <= to < len(order):
            item = order.pop(drag_start[0]); order.insert(to, item); refresh(to)
        drag_start[0] = None
    listing.bind("<Button-1>", on_press)
    listing.bind("<ButtonRelease-1>", on_release)
    refresh()
    result = [None]
    footer = ttk.Frame(window)
    footer.pack(fill="x", padx=18, pady=14)
    ttk.Button(footer, text=label(cfg, "إلغاء", "Cancel"), command=window.destroy).pack(side="left")
    def accept(): result[0] = order[:]; window.destroy()
    ttk.Button(footer, text=label(cfg, "إنشاء PDF", "Create PDF"), command=accept).pack(side="right")
    window.wait_window()
    return result[0]


def review_text(text: str, cfg: configparser.ConfigParser) -> str | None:
    if not cfg.getboolean("general", "review_text", fallback=True):
        return text
    window = modal(label(cfg, "مراجعة النص", "Review text"), 780, 580)
    ttk.Label(window, text=label(cfg, "يمكنك تصحيح النص قبل حفظه", "Correct the extracted text before saving"),
              font=("Segoe UI", 11)).pack(pady=10)
    box = tk.Text(window, wrap="word", undo=True, font=("Segoe UI", 11))
    box.pack(fill="both", expand=True, padx=16)
    box.insert("1.0", text)
    result = [None]
    footer = ttk.Frame(window)
    footer.pack(fill="x", padx=16, pady=12)
    ttk.Button(footer, text=label(cfg, "إلغاء", "Cancel"), command=window.destroy).pack(side="left")
    def accept(): result[0] = box.get("1.0", "end-1c"); window.destroy()
    ttk.Button(footer, text=label(cfg, "حفظ", "Save"), command=accept).pack(side="right")
    window.wait_window()
    return result[0]


def settings_ui(cfg: configparser.ConfigParser) -> None:
    window = modal(label(cfg, "إعدادات FileOrbit", "FileOrbit settings"), 520, 610)
    frame = ttk.Frame(window, padding=20)
    frame.pack(fill="both", expand=True)
    ttk.Label(frame, text=label(cfg, "الإعدادات", "Settings"), font=("Segoe UI", 16, "bold")).pack(anchor="w", pady=(0, 14))
    values = cfg["general"]
    language = tk.StringVar(value=values.get("language", "ar"))
    right = tk.BooleanVar(value=values.get("right_hold", "1") == "1")
    shift = tk.BooleanVar(value=values.get("shift_drag", "1") == "1")
    startup = tk.BooleanVar(value=values.get("startup", "1") == "1")
    review_ocr = tk.BooleanVar(value=values.get("review_text", "1") == "1")
    review_images = tk.BooleanVar(value=values.get("review_images", "1") == "1")
    sort = tk.StringVar(value=values.get("sort", "name"))
    descending = tk.BooleanVar(value=values.get("sort_desc", "0") == "1")
    hold = tk.StringVar(value=values.get("hold_ms", "500"))
    quality = tk.StringVar(value=values.get("quality", "balanced"))
    def row(text_ar, text_en, widget):
        line = ttk.Frame(frame); line.pack(fill="x", pady=6)
        ttk.Label(line, text=label(cfg, text_ar, text_en), width=29).pack(side="left")
        widget(line).pack(side="right")
    row("لغة الواجهة", "Language", lambda p: ttk.Combobox(p, values=["ar", "en"], textvariable=language, state="readonly", width=14))
    ttk.Checkbutton(frame, text=label(cfg, "ضغط أيمن مطوّل", "Hold right mouse button"), variable=right).pack(anchor="w", pady=5)
    ttk.Checkbutton(frame, text=label(cfg, "سحب مع Shift", "Shift + drag"), variable=shift).pack(anchor="w", pady=5)
    row("مدة الضغط (مللي ثانية)", "Hold duration (ms)", lambda p: ttk.Entry(p, textvariable=hold, width=17))
    ttk.Checkbutton(frame, text=label(cfg, "تشغيل عند بدء ويندوز", "Start with Windows"), variable=startup).pack(anchor="w", pady=5)
    ttk.Separator(frame).pack(fill="x", pady=10)
    ttk.Checkbutton(frame, text=label(cfg, "مراجعة النص قبل الحفظ", "Review OCR text before saving"), variable=review_ocr).pack(anchor="w", pady=5)
    ttk.Checkbutton(frame, text=label(cfg, "مراجعة ترتيب الصور", "Review image order"), variable=review_images).pack(anchor="w", pady=5)
    row("ترتيب الصور الافتراضي", "Default image order", lambda p: ttk.Combobox(p, values=["name", "time"], textvariable=sort, state="readonly", width=14))
    ttk.Checkbutton(frame, text=label(cfg, "من الأحدث إلى الأقدم / عكس الاسم", "Reverse order"), variable=descending).pack(anchor="w", pady=5)
    row("جودة الوسائط", "Media quality", lambda p: ttk.Combobox(p, values=["high", "balanced", "small"], textvariable=quality, state="readonly", width=14))
    ttk.Label(frame, text=label(cfg, "جميع المعالجات محلية ولا تُرسل الملفات للإنترنت.", "All processing is local; files are not uploaded."),
              foreground="#666666", wraplength=470).pack(anchor="w", pady=(20, 0))
    def save():
        try: milliseconds = max(250, min(1200, int(hold.get())))
        except ValueError: milliseconds = 500
        values.update({"language": language.get(), "right_hold": "1" if right.get() else "0",
                       "shift_drag": "1" if shift.get() else "0", "hold_ms": str(milliseconds),
                       "startup": "1" if startup.get() else "0", "review_text": "1" if review_ocr.get() else "0",
                       "review_images": "1" if review_images.get() else "0", "sort": sort.get(),
                       "sort_desc": "1" if descending.get() else "0", "quality": quality.get()})
        save_settings(cfg)
        window.destroy()
    ttk.Button(frame, text=label(cfg, "حفظ الإعدادات", "Save settings"), command=save).pack(anchor="e", pady=17)
    window.wait_window()


class ImageEditor:
    def __init__(self, path: Path, mode: str, cfg: configparser.ConfigParser):
        from PIL import Image, ImageOps, ImageTk
        self.path, self.mode, self.cfg = path, mode, cfg
        self.image = ImageOps.exif_transpose(Image.open(path)).convert("RGBA")
        self.window = modal(label(cfg, "تعديل الصورة", "Edit image"), 920, 670)
        self.result = None
        self.crop_box = None
        self.drag_start = None
        self.preview_box = (0, 0, 1, 1)
        self.preview_photo = None
        body = ttk.Frame(self.window, padding=12)
        body.pack(fill="both", expand=True)
        self.canvas = tk.Canvas(body, width=620, height=510, bg="#25232a", highlightthickness=0)
        self.canvas.pack(side="left", fill="both", expand=True)
        self.canvas.bind("<Button-1>", self.crop_start)
        self.canvas.bind("<B1-Motion>", self.crop_move)
        self.canvas.bind("<ButtonRelease-1>", self.crop_end)
        controls = ttk.Frame(body, width=250)
        controls.pack(side="right", fill="y", padx=(16, 0))
        self.rotation = tk.StringVar(value="0")
        self.resize = tk.StringVar(value="100")
        self.brightness = tk.DoubleVar(value=1.0)
        self.contrast = tk.DoubleVar(value=1.0)
        self.text = tk.StringVar()
        self.padding = tk.StringVar(value="80")
        self.radius = tk.StringVar(value="24")
        self.shadow = tk.StringVar(value="12")
        self.bg_color = tk.StringVar(value="#f3c989")
        self.remove_bg = tk.BooleanVar(value=False)
        def field(ar, en, var):
            ttk.Label(controls, text=label(cfg, ar, en)).pack(anchor="w", pady=(7, 1))
            ttk.Entry(controls, textvariable=var, width=20).pack(anchor="w")
        if mode in {"edit", "crop"}:
            ttk.Label(controls, text=label(cfg, "اسحب فوق الصورة لتحديد القص", "Drag on image to crop"), wraplength=230).pack(anchor="w")
        if mode == "edit":
            field("تدوير بالدرجات", "Rotation degrees", self.rotation)
            field("الحجم %", "Resize %", self.resize)
            ttk.Label(controls, text=label(cfg, "الإضاءة", "Brightness")).pack(anchor="w", pady=(7, 0))
            ttk.Scale(controls, from_=0.3, to=2.0, variable=self.brightness).pack(fill="x")
            ttk.Label(controls, text=label(cfg, "التباين", "Contrast")).pack(anchor="w", pady=(7, 0))
            ttk.Scale(controls, from_=0.3, to=2.0, variable=self.contrast).pack(fill="x")
            field("نص على الصورة", "Text overlay", self.text)
        if mode == "background":
            ttk.Checkbutton(controls, text=label(cfg, "إزالة الخلفية", "Remove background"), variable=self.remove_bg).pack(anchor="w", pady=8)
            field("لون الخلفية", "Background color", self.bg_color)
            field("المسافة حول الصورة", "Padding", self.padding)
            field("استدارة الزوايا", "Corner radius", self.radius)
            field("الظل", "Shadow", self.shadow)
            ttk.Button(controls, text=label(cfg, "اختيار لون", "Choose color"), command=self.choose_color).pack(anchor="w", pady=8)
        ttk.Button(controls, text=label(cfg, "حفظ نسخة جديدة", "Save new copy"), command=self.save).pack(anchor="w", pady=(22, 5))
        ttk.Button(controls, text=label(cfg, "إلغاء", "Cancel"), command=self.window.destroy).pack(anchor="w")
        self.window.after(100, self.render)
        self.window.wait_window()

    def choose_color(self):
        from tkinter.colorchooser import askcolor
        color = askcolor(color=self.bg_color.get(), parent=self.window)[1]
        if color: self.bg_color.set(color)

    def render(self):
        from PIL import ImageTk
        self.canvas.update_idletasks()
        width, height = max(400, self.canvas.winfo_width()), max(300, self.canvas.winfo_height())
        copy = self.image.copy()
        copy.thumbnail((width - 30, height - 30))
        self.preview_photo = ImageTk.PhotoImage(copy)
        x, y = (width - copy.width) // 2, (height - copy.height) // 2
        self.preview_box = (x, y, copy.width, copy.height)
        self.canvas.delete("all")
        self.canvas.create_image(x, y, anchor="nw", image=self.preview_photo)

    def crop_start(self, event):
        self.drag_start = (event.x, event.y)
        self.canvas.delete("crop")

    def crop_move(self, event):
        if self.drag_start is None: return
        self.canvas.delete("crop")
        self.canvas.create_rectangle(*self.drag_start, event.x, event.y, outline="#ffb347", width=3, tags="crop")

    def crop_end(self, event):
        if self.drag_start is None: return
        x, y, w, h = self.preview_box
        ax, ay = self.drag_start
        bx, by = event.x, event.y
        left = max(0, min(w, min(ax, bx) - x))
        top = max(0, min(h, min(ay, by) - y))
        right = max(0, min(w, max(ax, bx) - x))
        bottom = max(0, min(h, max(ay, by) - y))
        if right - left > 5 and bottom - top > 5:
            self.crop_box = tuple(round(value * scale) for value, scale in zip(
                (left, top, right, bottom),
                (self.image.width / w, self.image.height / h, self.image.width / w, self.image.height / h)))
        self.drag_start = None

    def save(self):
        from PIL import Image, ImageChops, ImageDraw, ImageEnhance, ImageFilter, ImageFont
        try:
            image = self.image.copy()
            if self.crop_box and self.mode in {"crop", "edit"}: image = image.crop(self.crop_box)
            if self.mode == "edit":
                angle = float(self.rotation.get() or "0")
                if angle: image = image.rotate(-angle, expand=True)
                scale = max(0.1, min(10, float(self.resize.get() or "100") / 100))
                if scale != 1: image = image.resize((max(1, round(image.width * scale)), max(1, round(image.height * scale))), Image.Resampling.LANCZOS)
                image = ImageEnhance.Brightness(image).enhance(self.brightness.get())
                image = ImageEnhance.Contrast(image).enhance(self.contrast.get())
                if self.text.get():
                    draw = ImageDraw.Draw(image)
                    font = ImageFont.truetype("arial.ttf", max(20, image.width // 25))
                    draw.text((20, 20), self.text.get(), font=font, fill="white", stroke_width=2, stroke_fill="black")
            if self.mode == "background":
                if self.remove_bg.get():
                    progress = tk.Toplevel(self.window); progress.title("FileOrbit")
                    ttk.Label(progress, text=label(self.cfg, "جاري إزالة الخلفية...", "Removing background...")).pack(padx=30, pady=25)
                    progress.update()
                    try:
                        image = remove_background_offline(image)
                    finally:
                        progress.destroy()
                else:
                    from PIL import ImageColor
                    pad = max(0, min(500, int(self.padding.get() or "0")))
                    corner = max(0, min(200, int(self.radius.get() or "0")))
                    shadow = max(0, min(80, int(self.shadow.get() or "0")))
                    new_size = (image.width + 2 * pad, image.height + 2 * pad)
                    backdrop = Image.new("RGBA", new_size, ImageColor.getrgb(self.bg_color.get()) + (255,))
                    if shadow:
                        mask = Image.new("L", new_size, 0)
                        ImageDraw.Draw(mask).rounded_rectangle((pad + 6, pad + 8, pad + image.width + 6, pad + image.height + 8), radius=corner, fill=150)
                        mask = mask.filter(ImageFilter.GaussianBlur(shadow))
                        dark = Image.new("RGBA", new_size, (25, 20, 30, 0))
                        dark.putalpha(mask)
                        backdrop.alpha_composite(dark)
                    if corner:
                        mask = Image.new("L", image.size, 0)
                        ImageDraw.Draw(mask).rounded_rectangle((0, 0, image.width, image.height), radius=corner, fill=255)
                        image.putalpha(ImageChops.multiply(image.getchannel("A"), mask))
                    backdrop.alpha_composite(image, (pad, pad))
                    image = backdrop
            suffix = "_cropped" if self.mode == "crop" else "_background" if self.mode == "background" else "_edited"
            out = unique_path(self.path.with_name(self.path.stem + suffix + ".png"))
            image.save(out, "PNG")
            self.result = out
            self.window.destroy()
        except Exception as error:
            messagebox.showerror("FileOrbit", str(error), parent=self.window)


def remove_background_offline(image):
    """Run the bundled U2NetP model locally without loading editing libraries at idle."""
    import numpy as np
    import onnxruntime as ort
    from PIL import Image
    model = APP_DIR / "assets" / "rembg" / "models" / "u2netp" / "u2netp.onnx"
    if not model.exists():
        raise RuntimeError("Background model is missing. / نموذج إزالة الخلفية غير موجود")
    sample = image.convert("RGB").resize((320, 320), Image.Resampling.LANCZOS)
    values = np.asarray(sample, dtype=np.float32)
    values /= max(float(values.max()), 1e-6)
    values = (values - np.array([.485, .456, .406], dtype=np.float32)) / np.array([.229, .224, .225], dtype=np.float32)
    tensor = np.transpose(values, (2, 0, 1))[np.newaxis, ...].astype(np.float32)
    session = ort.InferenceSession(str(model), providers=["CPUExecutionProvider"])
    prediction = session.run(None, {session.get_inputs()[0].name: tensor})[0][0, 0]
    minimum, maximum = float(prediction.min()), float(prediction.max())
    if maximum > minimum:
        prediction = (prediction - minimum) / (maximum - minimum)
    mask = Image.fromarray((np.clip(prediction, 0, 1) * 255).astype("uint8"), "L")
    mask = mask.resize(image.size, Image.Resampling.LANCZOS)
    return Image.composite(image.convert("RGBA"), Image.new("RGBA", image.size, 0), mask)


def split_pdf(path: Path) -> Path:
    import pymupdf as fitz
    folder = unique_dir(path.with_name(path.stem + "_pages_pdf"))
    folder.mkdir()
    try:
        with fitz.open(path) as source:
            for index in range(len(source)):
                with fitz.open() as target:
                    target.insert_pdf(source, from_page=index, to_page=index)
                    target.save(folder / f"{path.stem}_{index + 1:03d}.pdf")
        return folder
    except Exception:
        shutil.rmtree(folder, ignore_errors=True)
        raise


def merge_pdfs(files: list[Path]) -> Path:
    import pymupdf as fitz
    out = unique_path(files[0].with_name("Merged PDF.pdf"))
    with fitz.open() as target:
        for path in files:
            with fitz.open(path) as source: target.insert_pdf(source)
        target.save(out)
    return out


def act_on_file(path: Path, action: str, cfg: configparser.ConfigParser, root: tk.Tk) -> Path | None:
    kind = category(path)
    if action == "extract" and kind == "archive": return extract_archive(path)
    if action == "pages" and kind == "pdf": return split_pdf(path)
    if action in {"crop", "edit", "background"} and kind == "image":
        editor = ImageEditor(path, action, cfg)
        return editor.result
    if action == "to_pdf" and kind == "image": return pdf_from_images([path], path.with_suffix(".pdf"), cfg)
    if action == "ocr" and kind == "pdf":
        target = choose_one(cfg, "احفظ النص كـ", "Save text as", [("txt", "نص TXT", "TXT"), ("docx", "Word DOCX", "Word DOCX")])
        if target is None: return None
        text = review_text(pdf_text(path), cfg)
        return save_text_as(text, path, target) if text is not None else None
    if action == "compress" and kind == "pdf": return pdf_compress(path)
    if action == "compress" and kind == "image":
        from PIL import Image, ImageOps
        with Image.open(path) as source:
            image = ImageOps.exif_transpose(source)
            out = unique_path(path.with_name(path.stem + "_compressed.jpg"))
            image.convert("RGB").save(out, "JPEG", quality=68, optimize=True)
            return out
    if action in {"convert", "compress", "extract_audio"} and kind in {"video", "audio"}:
        if action == "extract_audio": target = "mp3"
        elif action == "compress": target = "mp4" if kind == "video" else "mp3"
        else:
            choices = [(x, x.upper(), x.upper()) for x in (["mp4", "mkv", "webm", "mov"] if kind == "video" else ["mp3", "wav", "flac", "m4a", "ogg"])]
            target = choose_one(cfg, "صيغة التحويل", "Output format", choices)
        if target is None: return None
        return run_media(path, target, action, cfg.get("general", "quality", fallback="balanced"), root)
    if action == "convert" and kind == "pdf":
        target = choose_one(cfg, "صيغة التحويل", "Output format", [("png", "صور PNG", "PNG images"), ("jpg", "صور JPG", "JPG images"), ("txt", "نص TXT", "TXT"), ("docx", "Word DOCX", "Word DOCX")])
        if target is None: return None
        if target in {"png", "jpg"}: return pdf_to_images(path, target)
        text = review_text(pdf_text(path), cfg)
        return save_text_as(text, path, target) if text is not None else None
    if action == "convert" and kind == "image":
        target = choose_one(cfg, "صيغة التحويل", "Output format", [(x, x.upper(), x.upper()) for x in ["jpg", "png", "webp", "bmp", "tiff", "pdf"]])
        return image_convert(path, target) if target else None
    if action == "convert" and kind == "document":
        target = choose_one(cfg, "صيغة التحويل", "Output format", [(x, x.upper(), x.upper()) for x in ["txt", "docx", "pdf"]])
        return document_convert(path, target) if target else None
    if action == "convert" and kind == "archive":
        target = choose_one(cfg, "صيغة الأرشيف", "Archive format", [(x, x.upper(), x.upper()) for x in ["zip", "7z", "tar", "tar.gz"]])
        if target is None: return None
        with tempfile.TemporaryDirectory() as temp:
            folder = extract_archive(path)
            try:
                contents = list(folder.iterdir())
                if not contents: raise RuntimeError("Archive is empty")
                output = create_archive(contents, target)
                final = unique_path(path.with_name(path.stem + "." + target))
                shutil.move(output, final)
                return final
            finally:
                shutil.rmtree(folder, ignore_errors=True)
    raise RuntimeError(f"No {action} operation for {path.name}")


def perform(job: dict, cfg: configparser.ConfigParser, root: tk.Tk) -> list[Path]:
    action = job["action"]
    if action == "settings":
        settings_ui(cfg); return []
    paths = [Path(p) for p in job.get("files", []) if Path(p).exists()]
    if not paths: raise RuntimeError("No accessible files were selected")
    if action == "archive":
        target = choose_one(cfg, "صيغة الأرشيف", "Archive format", [(x, x.upper(), x.upper()) for x in ["zip", "7z", "tar", "tar.gz"]], "zip")
        return [create_archive(paths, target)] if target else []
    if action == "to_pdf" and all(category(p) == "image" for p in paths):
        out = paths[0].with_name("Combined Images.pdf") if len(paths) > 1 else paths[0].with_suffix(".pdf")
        return [pdf_from_images(paths, out, cfg)]
    if action == "merge" and all(category(p) == "pdf" for p in paths):
        return [merge_pdfs(paths)]
    if action == "each":
        results = []
        for path in paths:
            kind = category(path)
            choices = {
                "image": [("convert", "تحويل", "Convert"), ("compress", "ضغط", "Compress"), ("to_pdf", "إلى PDF", "To PDF")],
                "pdf": [("convert", "تحويل", "Convert"), ("compress", "ضغط", "Compress"), ("ocr", "النص", "Text / OCR")],
                "video": [("convert", "تحويل", "Convert"), ("compress", "ضغط", "Compress"), ("extract_audio", "استخراج صوت", "Extract audio")],
                "audio": [("convert", "تحويل", "Convert"), ("compress", "ضغط", "Compress")],
                "archive": [("extract", "فك الأرشيف", "Extract archive"), ("convert", "إعادة أرشفة", "Repack")],
                "document": [("convert", "تحويل", "Convert")],
            }.get(kind, [("archive", "أرشفة", "Archive")])
            chosen = choose_one(cfg, f"عملية للملف {path.name}", f"Action for {path.name}",
                                choices)
            if chosen is None: continue
            try:
                result = create_archive([path], "zip") if chosen == "archive" else act_on_file(path, chosen, cfg, root)
                if result: results.append(result)
            except Exception as error:
                messagebox.showerror("FileOrbit", f"{path.name}: {error}", parent=root)
        return results
    results = []
    for path in paths:
        result = act_on_file(path, action, cfg, root)
        if result: results.append(result)
    return results


def main() -> int:
    global ROOT
    if len(sys.argv) != 3 or sys.argv[1] != "--job":
        print("Usage: FileOrbit.Worker.exe --job path.json", file=sys.stderr)
        return 2
    job_path = Path(sys.argv[2])
    job = json.loads(job_path.read_text(encoding="utf-8-sig"))
    job_path.unlink(missing_ok=True)
    cfg = load_settings()
    ROOT = tk.Tk()
    ROOT.withdraw()
    try:
        results = perform(job, cfg, ROOT)
        if results:
            names = "\n".join(str(p) for p in results[:8])
            if len(results) > 8: names += f"\n... +{len(results) - 8}"
            messagebox.showinfo("FileOrbit", label(cfg, "تم الحفظ:\n", "Saved:\n") + names, parent=ROOT)
        return 0
    except InterruptedError:
        return 0
    except Exception as error:
        messagebox.showerror("FileOrbit", str(error), parent=ROOT)
        return 1
    finally:
        ROOT.destroy()


if __name__ == "__main__":
    raise SystemExit(main())

